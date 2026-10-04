import openpyxl, re, json
import os
F=os.environ.get('MCR_WORKBOOK','Supplementary_Table_S1_per_genome_results.xlsx')
t=open('checkm2_isescan_raw.txt').read()
ck=[l.split('\t') for l in t.split('#CK')[1].split('#ISE')[0].strip().split('\n')]
ise=[l.split('\t') for l in t.split('#ISE')[1].strip().split('\n')]
# normalise checkm names -> galaxy input id
def norm(n):
    m=re.match(r'(GCA_\d+)[._](\d)_',n)
    if m: return f'{m.group(1)}.{m.group(2)}'
    return n.split('_')[0]
CK={norm(r[0]):(r[1],r[2]) for r in ck}
ISE={r[0]:r[4] for r in ise}
wb=openpyxl.load_workbook(F)
ws=wb['Isolates']
hdr={c.value:c.column for c in ws[1]}
gid_col=None
g=wb['Galaxy_results']; gh={c.value:c.column for c in g[1]}
bio2gid={g.cell(r,gh['BioSample']).value: g.cell(r,gh['Galaxy input ID']).value for r in range(2,g.max_row+1)}
c0=ws.max_column
ws.cell(1,c0+1,'CheckM2 completeness (%)')
ws.cell(1,c0+2,'CheckM2 contamination (%)')
ws.cell(1,c0+3,'IS30-family element on mcr contig')
hit=miss=0
for r in range(2,ws.max_row+1):
    bs=ws.cell(r,hdr['BioSample']).value
    if not bs: continue
    gid=bio2gid.get(bs)
    v=CK.get(str(gid))
    if v:
        ws.cell(r,c0+1,float(v[0])); ws.cell(r,c0+2,float(v[1])); hit+=1
    else: miss+=1
    iv=ISE.get(str(gid))
    if iv is not None: ws.cell(r,c0+3,'yes' if int(iv)>0 else 'no')
wb.save(F)
print('checkm filled',hit,'missing',miss)
vals=[float(v[0]) for v in CK.values()]; cont=[float(v[1]) for v in CK.values()]
print('completeness',min(vals),max(vals),'contamination',min(cont),max(cont))
print('n>=1% contam',sum(1 for x in cont if x>=1),'n>=5%',sum(1 for x in cont if x>=5))
