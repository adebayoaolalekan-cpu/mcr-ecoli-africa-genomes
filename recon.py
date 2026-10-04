import pandas as pd, re, collections
from openpyxl import load_workbook
import os
F=os.environ.get('MCR_WORKBOOK','Supplementary_Table_S1_per_genome_results.xlsx')
rows=[l.rstrip('\n').split('\t') for l in open('galaxy_typing.tsv') if l.strip()]
A=collections.defaultdict(list); P=collections.defaultdict(list); M={}; B={}
for r in rows:
    if r[0]=='A': A[r[1]].append(dict(contig=r[2],sym=r[3],meth=r[4],cov=float(r[5]),idt=float(r[6])))
    elif r[0]=='P': P[r[1]].append((r[2],r[3]))
    elif r[0]=='M': M[r[1]]=(r[2],r[3],r[4])
    elif r[0]=='B': B[(r[1],r[2])]=dict(mol=r[3],cl=r[4],reps=r[5],relax=r[6],binsize=int(r[7]),clen=int(r[8]))
Q={}; Fp={}
for l in open('galaxy_qc.tsv'):
    f=l.rstrip('\n').split('\t')
    if f[0]=='Q':
        n=f[1]; m=re.match(r'(GCA_\d+)_(\d)_',n); g=f'{m.group(1)}.{m.group(2)}' if m else n
        Q[g]=(int(f[2]),int(f[3]),int(f[4]),float(f[5]))
    else: Fp[f[1]]=int(f[2])
fam=lambda s: re.sub(r'\(.*?\)','',re.sub(r'_\d+(_\w+)?$','',s)).replace('IncHI2A','IncHI2')
def fams_from_text(s):
    return set(fam(x) for x in re.findall(r'(Inc(?:I-gamma/K1|B/O/K/Z|L/M|[A-Za-z0-9\-]+)|rep_cluster_\d+)',s or '') ) if s else set()
wb=load_workbook(F,data_only=True)
ws=wb['Isolates']; h={c.value:c.column for c in ws[1]}
g=wb['Galaxy_results']; gh={c.value:c.column for c in g[1]}
grow={g.cell(r,1).value:r for r in range(2,g.max_row+1)}
issues=[]; checked=collections.Counter()
for r in range(2,ws.max_row+1):
    bios=ws.cell(r,h['BioSample']).value
    cat=ws.cell(r,h['Recommended Galaxy input']).value
    gid = ws.cell(r,h['SRA Illumina PE run(s)']).value if cat=='SRA Illumina PE reads' else ws.cell(r,h['Assembly (GCA)']).value
    if gid=='SRR9099604': gid='GCA_006348965.1'
    st=ws.cell(r,h['ST (Achtman) - to fill']).value
    rep=ws.cell(r,h['Plasmid replicon carrying mcr - to fill']).value
    gr=grow[bios]
    # 1 Galaxy input id
    if g.cell(gr,gh['Galaxy input ID']).value!=gid: issues.append((bios,gid,'Galaxy input ID',g.cell(gr,gh['Galaxy input ID']).value,gid))
    # 2 ST
    m=M.get(gid); checked['ST']+=1
    if m is None: issues.append((bios,gid,'ST','missing in Galaxy',st))
    else:
        exp = int(m[1]) if m[1].isdigit() else 'Untypeable'
        if m[0]!='ecoli_achtman_4': issues.append((bios,gid,'MLST scheme',m[0],''))
        if st!=exp: issues.append((bios,gid,'ST (Isolates)',st,exp))
        if str(g.cell(gr,gh['MLST ST (raw)']).value)!=m[1]: issues.append((bios,gid,'MLST ST (raw)',g.cell(gr,gh['MLST ST (raw)']).value,m[1]))
        al=g.cell(gr,gh['MLST alleles']).value or ''
        nums=re.findall(r'\(([^)]*)\)',al) or al.split()
        if [x for x in nums]!=m[2].split(): issues.append((bios,gid,'MLST alleles',al,m[2]))
    # 3 mcr call
    calls=A.get(gid,[]); checked['mcr']+=1
    syms=list(dict.fromkeys(c['sym'] for c in calls))
    wsyms=[x for x in (g.cell(gr,gh['mcr (Galaxy AMRFinderPlus)']).value or '').split(';') if x]
    if set(syms)!=set(wsyms): issues.append((bios,gid,'mcr call',wsyms,syms))
    wcont=[x.strip() for x in (g.cell(gr,gh['mcr contig']).value or '').split(';') if x.strip()]
    if set(wcont)!=set(c['contig'] for c in calls): issues.append((bios,gid,'mcr contig',wcont,[c['contig'] for c in calls]))
    for k,col in (('cov','% coverage'),('idt','% identity')):
        wv=[float(x) for x in str(g.cell(gr,gh[col]).value or '').split(';') if x.strip()]
        sv=sorted(set(c[k] for c in calls)); 
        if sorted(set(wv))!=sv: issues.append((bios,gid,col,wv,sv))
    # 4 replicon
    checked['replicon']+=1
    on_contig=set(fam(p[1]) for c in calls for p in P.get(gid,[]) if p[0]==c['contig'])
    wf=fams_from_text(rep)
    if gid=='GCA_007993675.2':
        exp={'IncI2'}; src='decision (plasmid contig CP042471.2)'
        if not set(fam(p[1]) for p in P[gid] if p[0]=='CP042471.2')=={'IncI2'}: issues.append((bios,gid,'replicon source','',''))
    elif on_contig:
        exp=on_contig; src='PlasmidFinder same contig'
    else:
        b=[B.get((gid,c['contig'])) for c in calls]; b=[x for x in b if x]
        if not b: exp=set(); src='none'
        else:
            x=b[0]
            if x['mol']=='chromosome': exp={'CHROM'}; src='MOB chromosome'
            else:
                reps=[fam(y) for y in x['reps'].split(',') if y and y!='-']
                inc=[y for y in reps if not y.startswith('rep_cluster')]
                exp=set(inc) if inc else (set(reps) if reps else {'NOREP'}); src='MOB bin'
    if exp=={'CHROM'}: ok = rep and rep.startswith('Chromosomal')
    elif exp=={'NOREP'}: ok = rep and 'no replicon' in rep
    else: ok = wf==exp
    if not ok: issues.append((bios,gid,'replicon ('+src+')',rep,sorted(exp)))
    # MOB columns in Galaxy_results
    b=[B.get((gid,c['contig'])) for c in calls if B.get((gid,c['contig']))]
    if b and gid!='GCA_007993675.2':
        x=b[0]; checked['mob']+=1
        if g.cell(gr,gh['MOB-Recon molecule']).value!=x['mol']: issues.append((bios,gid,'MOB molecule',g.cell(gr,gh['MOB-Recon molecule']).value,x['mol']))
        if set((g.cell(gr,gh['MOB-Recon bin replicon(s)']).value or '-').split(','))!=set(x['reps'].split(',')): issues.append((bios,gid,'MOB reps',g.cell(gr,gh['MOB-Recon bin replicon(s)']).value,x['reps']))
        wb_bs=g.cell(gr,gh['MOB-Recon bin size (bp)']).value
        if x['mol']=='plasmid' and wb_bs!=x['binsize']: issues.append((bios,gid,'MOB bin size',wb_bs,x['binsize']))
    # 5 QC
    if gid in Q:
        checked['qc']+=1
        q=Q[gid]; wq=(g.cell(gr,gh['Contigs (QUAST)']).value,g.cell(gr,gh['Total length (bp)']).value,g.cell(gr,gh['N50 (bp)']).value,g.cell(gr,gh['GC (%)']).value)
        if tuple(wq)!=q: issues.append((bios,gid,'QUAST',wq,q))
        if not(4.5e6<=q[1]<=6.2e6 and q[0]<=500 and q[2]>=25000): issues.append((bios,gid,'QUAST threshold fail','',q))
    if cat=='SRA Illumina PE reads':
        run=ws.cell(r,h['SRA Illumina PE run(s)']).value
        wv=g.cell(gr,gh['Clean bases after fastp']).value
        if run!='SRR9099604':
            checked['fastp']+=1
            if wv!=Fp.get(run): issues.append((bios,run,'fastp clean bases',wv,Fp.get(run)))
            if Fp.get(run,0)<153e6: issues.append((bios,run,'depth below 30x','',Fp.get(run)))
    # 6 PD vs Galaxy mcr (informational)
pd.set_option('display.width',250); pd.set_option('display.max_colwidth',120)
df=pd.DataFrame(issues,columns=['BioSample','Galaxy ID','Field','Workbook','Galaxy source'])
print(dict(checked)); print(len(df)); print(df.to_string())
df.to_csv('issues.tsv',sep='\t',index=False)
# summary table counts check
ser=[ws.cell(r,h['Plasmid replicon carrying mcr - to fill']).value for r in range(2,ws.max_row+1)]
print('filled ST',sum(1 for r in range(2,ws.max_row+1) if ws.cell(r,h['ST (Achtman) - to fill']).value not in (None,'')),'filled rep',sum(1 for x in ser if x))
