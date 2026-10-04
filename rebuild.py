import re, collections
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
import os
F=os.environ.get('MCR_WORKBOOK','Supplementary_Table_S1_per_genome_results.xlsx')
rows=[l.rstrip('\n').split('\t') for l in open('galaxy_typing.tsv') if l.strip()]
A=collections.defaultdict(list); P=collections.defaultdict(list); M={}; B={}
for r in rows:
    if r[0]=='A': A[r[1]].append(r[2:])
    elif r[0]=='P': P[r[1]].append((r[2],r[3]))
    elif r[0]=='M': M[r[1]]=(r[3],r[4])
    elif r[0]=='B': B[(r[1],r[2])]=r[3:]
Q={}; Fp={}
for l in open('galaxy_qc.tsv'):
    f=l.rstrip('\n').split('\t')
    if f[0]=='Q':
        m=re.match(r'(GCA_\d+)_(\d)_',f[1]); g=f'{m.group(1)}.{m.group(2)}' if m else f[1]
        Q[g]=(int(f[2]),int(f[3]),int(f[4]),float(f[5]))
    else: Fp[f[1]]=int(f[2])
# ABRicate NCBI/ResFinder mcr from earlier per-dataset pulls (compactA/B) + GCA_006348965.1 not run
R=collections.defaultdict(lambda: collections.defaultdict(list))
for fn in ('../compactB.tsv','../compactA.tsv'):
    for l in open(fn):
        f=l.rstrip('\n').split('\t')
        if f[0]=='R': R[f[2]][f[1]].append(f[4])
wb=load_workbook(F)
ws=wb['Isolates']; h={c.value:c.column for c in ws[1]}
alle=['adk','fumC','gyrB','icd','mdh','purA','recA']
cols=['BioSample','Galaxy input ID','Galaxy arm','Clean bases after fastp','Est. depth (x, 5.1 Mb genome)','Contigs (QUAST)','Total length (bp)','N50 (bp)','GC (%)','QC status',
      'mcr (PD)','mcr (Galaxy AMRFinderPlus)','AMRFinderPlus method','% coverage','% identity','mcr contig',
      'ABRicate NCBI mcr','ABRicate ResFinder mcr','PlasmidFinder replicon(s) on mcr contig','MOB-Recon molecule','MOB-Recon bin replicon(s)','MOB-Recon relaxase(s)','MOB-Recon bin size (bp)',
      'MLST ST (raw)','MLST alleles','All PlasmidFinder replicons in genome']
out=[]
for r in range(2,ws.max_row+1):
    bios=ws.cell(r,h['BioSample']).value; cat=ws.cell(r,h['Recommended Galaxy input']).value
    run=ws.cell(r,h['SRA Illumina PE run(s)']).value
    gid = run if cat=='SRA Illumina PE reads' else ws.cell(r,h['Assembly (GCA)']).value
    arm = 'A (reads, Shovill)' if cat=='SRA Illumina PE reads' else ('B (ENA assembly)' if gid=='GCA_988286365.1' else 'B (GenBank assembly)')
    clean = Fp.get(run) if cat=='SRA Illumina PE reads' else None
    if gid=='SRR9099604': gid='GCA_006348965.1'; arm='B (GenBank assembly; reads SRR9099604 failed depth QC)'
    q=Q.get(gid)
    qcs = 'Pass' if q and (4.5e6<=q[1]<=6.2e6 and q[0]<=500 and q[2]>=25000) else ('Not run (ENA upload; 10 contigs, 5.36 Mb)' if not q else 'Fail')
    calls=A.get(gid,[])
    uc=[]; seen=set()
    for c in calls:
        if (c[0],c[1]) in seen: continue
        seen.add((c[0],c[1])); uc.append(c)
    on=[f'{p[1]}' for c in uc for p in P.get(gid,[]) if p[0]==c[0]]
    b=next((B[(gid,c[0])] for c in uc if (gid,c[0]) in B),None)
    if gid=='GCA_007993675.2': b=None
    m=M.get(gid,('',''))
    al=' '.join(f'{a}({v})' for a,v in zip(alle,m[1].split())) if m[1] else ''
    out.append([bios,gid,arm,clean,None,*(q if q else (None,None,None,None)),qcs,
        ws.cell(r,h['mcr variant(s)']).value,'; '.join(dict.fromkeys(c[1] for c in uc)),'; '.join(c[2] for c in uc),'; '.join(c[3] for c in uc),'; '.join(c[4] for c in uc),'; '.join(dict.fromkeys(c[0] for c in uc)),
        ';'.join(dict.fromkeys(R[gid].get('ncbi',[]))) or ('not run' if gid=='GCA_006348965.1' else 'none'),
        ';'.join(dict.fromkeys(R[gid].get('resfinder',[]))) or ('not run' if gid=='GCA_006348965.1' else 'none'),
        ' / '.join(dict.fromkeys(on)) or 'none',
        (b[0] if b else ('plasmid (CP042471.2) and chromosome-size contig (CP042470.1)' if gid=='GCA_007993675.2' else '')),
        (b[2] if b else ''),(b[3] if b else ''),(int(b[4]) if b and b[0]=='plasmid' else None),
        m[0], al, ', '.join(sorted(set(p[1] for p in P.get(gid,[]))))])
del wb['Galaxy_results']
ns=wb.create_sheet('Galaxy_results')
hfont=Font(name='Arial',bold=True,color='FFFFFFFF'); hfill=PatternFill('solid',fgColor='FF1F4E78'); bfont=Font(name='Arial',size=10)
for j,c in enumerate(cols,1):
    x=ns.cell(1,j,c); x.font=hfont; x.fill=hfill; x.alignment=Alignment(wrap_text=True,vertical='top')
assert all(len(o)==len(cols) for o in out), set(len(o) for o in out)
for i,row in enumerate(out,2):
    for j,v in enumerate(row,1):
        if j==5: v=f'=IF(D{i}="","",D{i}/5100000)'
        ns.cell(i,j,v).font=bfont
    ns.cell(i,5).number_format='0.0'
    for j in (4,6,7,8,23): ns.cell(i,j).number_format='#,##0'
for j in range(1,len(cols)+1): ns.column_dimensions[get_column_letter(j)].width=16
for j in (3,19,21,25,26): ns.column_dimensions[get_column_letter(j)].width=32
ns.freeze_panes='C2'; ns.auto_filter.ref=f'A1:{get_column_letter(len(cols))}{len(out)+1}'; ns.row_dimensions[1].height=45
wb.save(F); print('rebuilt',len(out))
